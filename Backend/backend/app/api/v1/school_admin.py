from fastapi import APIRouter, Depends, HTTPException, status, File, UploadFile, Response
from sqlalchemy.orm import Session
from typing import List
from datetime import datetime, timezone

from app.core.database import get_db
from app.api.deps import require_roles, get_current_school_id, get_current_user
from app.models.user import User, UserRole
from app.models.lms import Grade, ParentStudentLink
from app.models.course import Course, CourseTeacher
from app.models.platform import ActivityLog
from app.core.security import get_password_hash, generate_temp_password
from app.schemas.user import UserCreate, UserResponse, UserUpdate, ParentCreate, ParentUpdate, ParentResponse
from app.schemas.lms import GradeCreate, GradeResponse
from app.schemas.course import CourseResponse
from app.schemas.sso import SSOConfigUpdate, SSOConfigResponse
from app.schemas.complaint import ComplaintCreate, ComplaintResponse
from app.models.complaint import Complaint

router = APIRouter(dependencies=[Depends(require_roles([UserRole.ADMIN]))])


# There's no self-service "forgot password" flow in this app (no email
# service is wired up to deliver a reset link -- see the notification
# system's same limitation), so password recovery is admin-initiated
# instead: School Admin resets a teacher/student/parent's password here and
# hands them the new one, exactly like account creation already does (same
# generate_temp_password + a one-time-reveal response). Restricted to
# TEACHER/STUDENT/PARENT -- an admin resetting another admin's password (or
# a super admin's) is deliberately out of scope here.
RESETTABLE_ROLES = [UserRole.TEACHER, UserRole.STUDENT, UserRole.PARENT]


@router.post("/users/{user_id}/reset-password", response_model=UserResponse)
def reset_user_password(
    user_id: str,
    school_id: str = Depends(get_current_school_id),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    user = (
        db.query(User)
        .filter(User.id == user_id, User.school_id == school_id, User.role.in_(RESETTABLE_ROLES))
        .first()
    )
    if not user:
        raise HTTPException(status_code=404, detail="User not found or access denied")

    new_password = generate_temp_password(user.full_name, fallback=user.role.value.title())
    user.hashed_password = get_password_hash(new_password)

    log = ActivityLog(
        school_id=school_id, user_id=current_user.id, user_name=current_user.full_name,
        action="PASSWORD_RESET", details=f"Reset password for {user.role.value.title()}: {user.full_name}",
    )
    db.add(log)

    db.commit()
    db.refresh(user)
    user.generated_password = new_password
    return user


@router.get("/teachers", response_model=List[UserResponse])
def get_teachers(school_id: str = Depends(get_current_school_id), db: Session = Depends(get_db)):
    return db.query(User).filter(User.school_id == school_id, User.role == UserRole.TEACHER).all()

@router.post("/teachers", response_model=UserResponse)
def create_teacher(
    payload: UserCreate,
    school_id: str = Depends(get_current_school_id),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    # Generate a secure password server-side (same approach used platform-wide --
    # see generate_temp_password) when the caller omits one, instead of requiring
    # a real login credential to be produced in the browser.
    generated_password = None
    raw_password = payload.password.strip() if payload.password and payload.password.strip() else None
    if not raw_password:
        raw_password = generate_temp_password(payload.full_name, fallback="Teacher")
        generated_password = raw_password

    teacher = User(
        email=payload.email,
        hashed_password=get_password_hash(raw_password),
        full_name=payload.full_name,
        role=UserRole.TEACHER,
        school_id=school_id
    )
    db.add(teacher)
    
    log = ActivityLog(school_id=school_id, user_id=current_user.id, user_name=current_user.full_name, action="TEACHER_CREATED", details=f"Teacher: {teacher.full_name}")
    db.add(log)

    db.commit()
    db.refresh(teacher)
    if generated_password:
        teacher.generated_password = generated_password
    return teacher

@router.put("/teachers/{teacher_id}", response_model=UserResponse)
def update_teacher(
    teacher_id: str,
    payload: UserUpdate,
    school_id: str = Depends(get_current_school_id),
    db: Session = Depends(get_db)
):
    teacher = db.query(User).filter(User.id == teacher_id, User.school_id == school_id, User.role == UserRole.TEACHER).first()
    if not teacher:
        raise HTTPException(status_code=404, detail="Teacher not found or access denied")
    if payload.full_name:
        teacher.full_name = payload.full_name
    if payload.email:
        teacher.email = payload.email
    if payload.is_active is not None:
        teacher.is_active = payload.is_active
    db.commit()
    db.refresh(teacher)
    return teacher

@router.delete("/teachers/{teacher_id}")
def delete_teacher(
    teacher_id: str,
    school_id: str = Depends(get_current_school_id),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    teacher = db.query(User).filter(User.id == teacher_id, User.school_id == school_id, User.role == UserRole.TEACHER).first()
    if not teacher:
        raise HTTPException(status_code=404, detail="Teacher not found or access denied")
    deleted_name = teacher.full_name
    db.delete(teacher)
    log = ActivityLog(school_id=school_id, user_id=current_user.id, user_name=current_user.full_name, action="TEACHER_DELETED", details=f"Teacher: {deleted_name}")
    db.add(log)
    db.commit()
    return {"message": "Teacher deleted", "id": teacher_id}


@router.get("/students", response_model=List[UserResponse])
def get_students(school_id: str = Depends(get_current_school_id), db: Session = Depends(get_db)):
    return db.query(User).filter(User.school_id == school_id, User.role == UserRole.STUDENT).all()

@router.post("/students", response_model=UserResponse)
def create_student(
    payload: UserCreate,
    school_id: str = Depends(get_current_school_id),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    # Generate a secure password server-side (same approach used platform-wide --
    # see generate_temp_password) when the caller omits one, instead of requiring
    # a real login credential to be produced in the browser.
    generated_password = None
    raw_password = payload.password.strip() if payload.password and payload.password.strip() else None
    if not raw_password:
        raw_password = generate_temp_password(payload.full_name, fallback="Student")
        generated_password = raw_password

    student = User(
        email=payload.email,
        hashed_password=get_password_hash(raw_password),
        full_name=payload.full_name,
        role=UserRole.STUDENT,
        grade_id=payload.grade_id,
        school_id=school_id
    )
    db.add(student)

    log = ActivityLog(school_id=school_id, user_id=current_user.id, user_name=current_user.full_name, action="STUDENT_CREATED", details=f"Student: {student.full_name}")
    db.add(log)

    db.commit()
    db.refresh(student)
    if generated_password:
        student.generated_password = generated_password
    return student

@router.put("/students/{student_id}", response_model=UserResponse)
def update_student(
    student_id: str,
    payload: UserUpdate,
    school_id: str = Depends(get_current_school_id),
    db: Session = Depends(get_db)
):
    student = db.query(User).filter(User.id == student_id, User.school_id == school_id, User.role == UserRole.STUDENT).first()
    if not student:
        raise HTTPException(status_code=404, detail="Student not found or access denied")
    if payload.full_name:
        student.full_name = payload.full_name
    if payload.email:
        student.email = payload.email
    if payload.grade_id:
        student.grade_id = payload.grade_id
    if payload.is_active is not None:
        student.is_active = payload.is_active
    db.commit()
    db.refresh(student)
    return student

def _delete_student_and_log(db: Session, student: User, school_id: str, current_user: User) -> str:
    """Shared hard-delete path: the plain admin-initiated delete below, and
    fulfilling a Task #61 deletion request, both end up here so the audit
    trail (and the cascade this relies on) only has one implementation."""
    deleted_name = student.full_name
    db.delete(student)
    log = ActivityLog(school_id=school_id, user_id=current_user.id, user_name=current_user.full_name, action="STUDENT_DELETED", details=f"Student: {deleted_name}")
    db.add(log)
    db.commit()
    return deleted_name


@router.delete("/students/{student_id}")
def delete_student(
    student_id: str,
    school_id: str = Depends(get_current_school_id),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    student = db.query(User).filter(User.id == student_id, User.school_id == school_id, User.role == UserRole.STUDENT).first()
    if not student:
        raise HTTPException(status_code=404, detail="Student not found or access denied")
    _delete_student_and_log(db, student, school_id, current_user)
    return {"message": "Student deleted", "id": student_id}


@router.get("/students/{student_id}/export-data")
def export_student_data_admin(
    student_id: str,
    school_id: str = Depends(get_current_school_id),
    db: Session = Depends(get_db),
):
    """Task #61: lets a School Admin pull a full personal-data export for one
    student -- e.g. to hand over before fulfilling a deletion request, or to
    answer a records request directly."""
    import json as _json
    from app.services.data_export_service import export_student_data

    student = db.query(User).filter(User.id == student_id, User.school_id == school_id, User.role == UserRole.STUDENT).first()
    if not student:
        raise HTTPException(status_code=404, detail="Student not found or access denied")

    data = export_student_data(db, student)
    return Response(
        content=_json.dumps(data, indent=2, default=str),
        media_type="application/json",
        headers={"Content-Disposition": f"attachment; filename={student.full_name.replace(' ', '_')}-data-export.json"},
    )


@router.get("/data-requests")
def list_data_deletion_requests(
    status_filter: str = None,
    school_id: str = Depends(get_current_school_id),
    db: Session = Depends(get_db),
):
    """Task #61: the School Admin's queue of student/parent-submitted
    deletion requests, defaulting to every request (newest first) so a
    pending badge count and a resolved history can both read from this
    one endpoint."""
    from app.models.data_request import DataDeletionRequest

    query = db.query(DataDeletionRequest).filter(DataDeletionRequest.school_id == school_id)
    if status_filter:
        query = query.filter(DataDeletionRequest.status == status_filter)
    requests = query.order_by(DataDeletionRequest.requested_at.desc()).all()
    return [
        {
            "id": r.id,
            "student_id": r.student_id,
            "student_name": r.student_name,
            "requested_by_role": r.requested_by_role,
            "status": r.status,
            "note": r.note,
            "requested_at": r.requested_at,
            "resolved_at": r.resolved_at,
        }
        for r in requests
    ]


@router.post("/data-requests/{request_id}/fulfill")
def fulfill_data_deletion_request(
    request_id: str,
    school_id: str = Depends(get_current_school_id),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Task #61: approving a deletion request actually erases the student --
    same hard-delete path as the plain DELETE /students/{id} route -- and
    marks the request fulfilled. A request whose student was already removed
    some other way is marked fulfilled without erroring."""
    from app.models.data_request import DataDeletionRequest

    request_row = db.query(DataDeletionRequest).filter(
        DataDeletionRequest.id == request_id, DataDeletionRequest.school_id == school_id
    ).first()
    if not request_row:
        raise HTTPException(status_code=404, detail="Request not found")
    if request_row.status != "pending":
        raise HTTPException(status_code=400, detail="This request has already been resolved")

    student = db.query(User).filter(
        User.id == request_row.student_id, User.school_id == school_id, User.role == UserRole.STUDENT
    ).first()
    if student:
        _delete_student_and_log(db, student, school_id, current_user)

    request_row.status = "fulfilled"
    request_row.resolved_at = datetime.now(timezone.utc)
    request_row.resolved_by_id = current_user.id
    db.commit()
    return {"message": "Deletion request fulfilled -- the student's account and data have been erased.", "id": request_id}


@router.post("/data-requests/{request_id}/dismiss")
def dismiss_data_deletion_request(
    request_id: str,
    school_id: str = Depends(get_current_school_id),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Task #61: closes a request without deleting anything -- e.g. it was
    submitted by mistake, or was resolved outside the system."""
    from app.models.data_request import DataDeletionRequest

    request_row = db.query(DataDeletionRequest).filter(
        DataDeletionRequest.id == request_id, DataDeletionRequest.school_id == school_id
    ).first()
    if not request_row:
        raise HTTPException(status_code=404, detail="Request not found")
    if request_row.status != "pending":
        raise HTTPException(status_code=400, detail="This request has already been resolved")

    request_row.status = "dismissed"
    request_row.resolved_at = datetime.now(timezone.utc)
    request_row.resolved_by_id = current_user.id
    db.commit()
    return {"message": "Request dismissed.", "id": request_id}


def _serialize_parent(db: Session, parent: User) -> dict:
    links = db.query(ParentStudentLink).filter(ParentStudentLink.parent_id == parent.id).all()
    children = []
    for link in links:
        student = db.query(User).filter(User.id == link.student_id).first()
        if student:
            children.append({"id": student.id, "name": student.full_name, "email": student.email})
    return {
        "id": parent.id,
        "email": parent.email,
        "full_name": parent.full_name,
        "is_active": parent.is_active,
        "generated_password": getattr(parent, "generated_password", None),
        "children": children,
    }


@router.get("/parents", response_model=List[ParentResponse])
def get_parents(school_id: str = Depends(get_current_school_id), db: Session = Depends(get_db)):
    parents = db.query(User).filter(User.school_id == school_id, User.role == UserRole.PARENT).all()
    return [_serialize_parent(db, p) for p in parents]


@router.post("/parents", response_model=ParentResponse)
def create_parent(
    payload: ParentCreate,
    school_id: str = Depends(get_current_school_id),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    generated_password = None
    raw_password = payload.password.strip() if payload.password and payload.password.strip() else None
    if not raw_password:
        raw_password = generate_temp_password(payload.full_name, fallback="Parent")
        generated_password = raw_password

    parent = User(
        email=payload.email,
        hashed_password=get_password_hash(raw_password),
        full_name=payload.full_name,
        role=UserRole.PARENT,
        school_id=school_id
    )
    db.add(parent)
    db.flush()  # assign parent.id before creating links that reference it

    # Only link students that actually belong to this school -- never trust
    # a student_id from the request body without checking tenant scope.
    valid_student_ids = {
        row[0] for row in db.query(User.id).filter(
            User.id.in_(payload.student_ids or []),
            User.school_id == school_id,
            User.role == UserRole.STUDENT,
        ).all()
    }
    for student_id in valid_student_ids:
        db.add(ParentStudentLink(school_id=school_id, parent_id=parent.id, student_id=student_id))

    log = ActivityLog(school_id=school_id, user_id=current_user.id, user_name=current_user.full_name, action="PARENT_CREATED", details=f"Parent: {parent.full_name}")
    db.add(log)

    db.commit()
    db.refresh(parent)
    if generated_password:
        parent.generated_password = generated_password
    return _serialize_parent(db, parent)


@router.put("/parents/{parent_id}", response_model=ParentResponse)
def update_parent(
    parent_id: str,
    payload: ParentUpdate,
    school_id: str = Depends(get_current_school_id),
    db: Session = Depends(get_db)
):
    parent = db.query(User).filter(parent_id == User.id, User.school_id == school_id, User.role == UserRole.PARENT).first()
    if not parent:
        raise HTTPException(status_code=404, detail="Parent not found or access denied")
    if payload.full_name:
        parent.full_name = payload.full_name
    if payload.email:
        parent.email = payload.email
    if payload.is_active is not None:
        parent.is_active = payload.is_active

    if payload.student_ids is not None:
        # Replace the full linked-children set with what was provided.
        db.query(ParentStudentLink).filter(ParentStudentLink.parent_id == parent.id).delete()
        valid_student_ids = {
            row[0] for row in db.query(User.id).filter(
                User.id.in_(payload.student_ids),
                User.school_id == school_id,
                User.role == UserRole.STUDENT,
            ).all()
        }
        for student_id in valid_student_ids:
            db.add(ParentStudentLink(school_id=school_id, parent_id=parent.id, student_id=student_id))

    db.commit()
    db.refresh(parent)
    return _serialize_parent(db, parent)


@router.delete("/parents/{parent_id}")
def delete_parent(
    parent_id: str,
    school_id: str = Depends(get_current_school_id),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    parent = db.query(User).filter(User.id == parent_id, User.school_id == school_id, User.role == UserRole.PARENT).first()
    if not parent:
        raise HTTPException(status_code=404, detail="Parent not found or access denied")
    deleted_name = parent.full_name
    db.delete(parent)
    log = ActivityLog(school_id=school_id, user_id=current_user.id, user_name=current_user.full_name, action="PARENT_DELETED", details=f"Parent: {deleted_name}")
    db.add(log)
    db.commit()
    return {"message": "Parent deleted", "id": parent_id}


@router.get("/grades", response_model=List[GradeResponse])
def get_grades(school_id: str = Depends(get_current_school_id), db: Session = Depends(get_db)):
    """
    Students/teachers/sections were previously always 0 on the Grades page --
    GradeResponse had no such fields, so the frontend's `g.students_count`
    etc. always fell back to their zero defaults. These are computed for
    real here: students_count/sections come straight off the students
    enrolled in the grade (User.grade_id/User.section); teachers_count is
    the distinct set of teachers (primary + co-teachers) on courses whose
    grade_level names this grade, since teachers aren't assigned to a
    grade directly in this schema.
    """
    grades = db.query(Grade).filter(Grade.school_id == school_id).all()

    for grade in grades:
        students = db.query(User).filter(
            User.school_id == school_id, User.role == UserRole.STUDENT, User.grade_id == grade.id
        ).all()
        grade.students_count = len(students)
        grade.sections = len({s.section for s in students if s.section})

        course_ids = [
            c.id for c in db.query(Course.id, Course.created_by_id).filter(
                Course.school_id == school_id, Course.grade_level == grade.name
            ).all()
        ]
        teacher_ids = {
            cid for (cid,) in db.query(Course.created_by_id).filter(
                Course.school_id == school_id, Course.grade_level == grade.name, Course.created_by_id.isnot(None)
            ).all()
        }
        if course_ids:
            teacher_ids |= {
                tid for (tid,) in db.query(CourseTeacher.teacher_id).filter(
                    CourseTeacher.course_id.in_(course_ids)
                ).all()
            }
        grade.teachers_count = len(teacher_ids)

    return grades

@router.post("/grades", response_model=GradeResponse)
def create_grade(
    payload: GradeCreate,
    school_id: str = Depends(get_current_school_id),
    db: Session = Depends(get_db)
):
    grade = Grade(school_id=school_id, name=payload.name, code=payload.code)
    db.add(grade)
    db.commit()
    db.refresh(grade)
    return grade

@router.delete("/grades/{grade_id}")
def delete_grade(
    grade_id: str,
    school_id: str = Depends(get_current_school_id),
    db: Session = Depends(get_db)
):
    grade = db.query(Grade).filter(Grade.id == grade_id, Grade.school_id == school_id).first()
    if not grade:
        raise HTTPException(status_code=404, detail="Grade not found or access denied")
    db.delete(grade)
    db.commit()
    return {"message": "Grade deleted", "id": grade_id}


@router.get("/activity-logs")
def get_school_activity_logs(school_id: str = Depends(get_current_school_id), db: Session = Depends(get_db)):
    return db.query(ActivityLog).filter(ActivityLog.school_id == school_id).order_by(ActivityLog.timestamp.desc()).all()


@router.get("/students/excel-template")
def download_student_template():
    """
    Downloads sample Excel (.xlsx) template for bulk student upload
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    import io
    from fastapi.responses import StreamingResponse

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Students Template"

    headers = ["Name", "Grade", "Section"]
    ws.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Clearly-labeled placeholder rows (not real people) -- the admin deletes/overwrites
    # these before filling in and uploading their actual student roster.
    sample_rows = [
        ["Example Student 1", "Grade 6", "A"],
        ["Example Student 2", "Grade 6", "A"],
        ["Example Student 3", "Grade 7", "B"],
        ["Example Student 4", "Grade 8", "A"],
    ]
    for row in sample_rows:
        ws.append(row)

    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 15

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=student_upload_template.xlsx"}
    )


@router.post("/students/upload-excel")
async def upload_students_excel(
    file: UploadFile = File(...),
    school_id: str = Depends(get_current_school_id),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    Bulk uploads students from .xlsx Excel file, validates rows, auto-generates secure credentials, and creates students.
    """
    import openpyxl
    import io
    import re
    import random
    from app.models.lms import StudentStats
    from app.models.school import School

    if not file.filename.endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Only .xlsx Excel files are supported")

    contents = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
        ws = wb.active
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to parse Excel file: {str(e)}")

    school = db.query(School).filter(School.id == school_id).first()
    school_domain = school.domain if (school and school.domain) else "school.edu"

    # Fetch existing users to prevent duplicate email collisions
    existing_users = db.query(User.email).all()
    existing_emails = {u[0].lower() for u in existing_users}

    # Fetch existing grades map in this school
    grades = db.query(Grade).filter(Grade.school_id == school_id).all()
    grade_map = {g.name.lower().strip(): g for g in grades}
    for g in grades:
        if g.code:
            grade_map[g.code.lower().strip()] = g

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(status_code=400, detail="Uploaded Excel file is empty")

    header_row = [str(cell).strip().lower() if cell is not None else "" for cell in rows[0]]
    
    # Identify column indices
    name_idx = -1
    grade_idx = -1
    section_idx = -1

    for idx, col_name in enumerate(header_row):
        if "name" in col_name or "student" in col_name:
            name_idx = idx
        elif "grade" in col_name or "class" in col_name:
            grade_idx = idx
        elif "section" in col_name:
            section_idx = idx

    if name_idx == -1 or grade_idx == -1:
        raise HTTPException(
            status_code=400,
            detail="Excel file must contain 'Name' and 'Grade' columns. 'Section' is optional."
        )

    created_count = 0
    skipped_count = 0
    duplicate_count = 0
    invalid_count = 0
    errors = []
    credentials_list = []

    for row_num, row in enumerate(rows[1:], start=2):
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue

        raw_name = str(row[name_idx]).strip() if name_idx < len(row) and row[name_idx] is not None else ""
        raw_grade = str(row[grade_idx]).strip() if grade_idx < len(row) and row[grade_idx] is not None else ""
        raw_section = str(row[section_idx]).strip() if (section_idx != -1 and section_idx < len(row) and row[section_idx] is not None) else ""

        if not raw_name:
            invalid_count += 1
            errors.append(f"Row {row_num}: Student name is missing")
            continue

        if not raw_grade:
            invalid_count += 1
            errors.append(f"Row {row_num}: Grade is missing for student '{raw_name}'")
            continue

        # Match or create Grade
        norm_grade = raw_grade.lower().strip()
        matched_grade = grade_map.get(norm_grade)
        if not matched_grade:
            # Create grade automatically if it doesn't exist
            matched_grade = Grade(school_id=school_id, name=raw_grade, code=raw_grade)
            db.add(matched_grade)
            db.commit()
            db.refresh(matched_grade)
            grade_map[norm_grade] = matched_grade

        # Generate username & password
        clean_parts = re.sub(r'[^a-zA-Z0-9\s]', '', raw_name).lower().split()
        base_user = ".".join(clean_parts) if clean_parts else f"student{random.randint(100, 999)}"
        username = f"{base_user}@{school_domain}"
        counter = 1
        while username in existing_emails:
            username = f"{base_user}{counter}@{school_domain}"
            counter += 1
        existing_emails.add(username)

        raw_password = generate_temp_password(raw_name, fallback="Student")

        new_student = User(
            email=username,
            hashed_password=get_password_hash(raw_password),
            full_name=raw_name,
            role=UserRole.STUDENT,
            school_id=school_id,
            grade_id=matched_grade.id,
            section=raw_section or "A"
        )
        db.add(new_student)
        db.commit()
        db.refresh(new_student)

        # Initialize student stats
        stats = StudentStats(school_id=school_id, student_id=new_student.id, xp=0, streak_days=0, badges=[])
        db.add(stats)
        db.commit()

        created_count += 1
        credentials_list.append({
            "id": new_student.id,
            "name": raw_name,
            "username": username,
            "email": username,
            "initialPassword": raw_password,
            "grade": matched_grade.name,
            "section": raw_section or "A"
        })

    # Log bulk upload activity
    log = ActivityLog(
        school_id=school_id,
        user_id=current_user.id,
        user_name=current_user.full_name,
        action="BULK_STUDENT_UPLOAD",
        details=f"Bulk imported {created_count} students from Excel file '{file.filename}'"
    )
    db.add(log)
    db.commit()

    return {
        "message": f"Successfully onboarded {created_count} students.",
        "createdCount": created_count,
        "skippedCount": skipped_count,
        "duplicateCount": duplicate_count,
        "invalidCount": invalid_count,
        "errors": errors,
        "credentials": credentials_list
    }


@router.put("/settings/school")
def update_school_settings(
    payload: dict,
    school_id: str = Depends(get_current_school_id),
    db: Session = Depends(get_db)
):
    from app.models.school import School
    school = db.query(School).filter(School.id == school_id).first()
    if not school:
        raise HTTPException(status_code=404, detail="School not found")

    if "name" in payload and payload["name"]:
        school.name = payload["name"]
    if "domain" in payload and payload["domain"]:
        school.domain = payload["domain"]
    
    db.commit()
    db.refresh(school)
    return {"message": "School settings updated successfully", "school": {"id": school.id, "name": school.name, "domain": school.domain}}


@router.put("/settings/profile")
def update_admin_profile(
    payload: dict,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    if "fullName" in payload and payload["fullName"]:
        current_user.full_name = payload["fullName"]
    elif "full_name" in payload and payload["full_name"]:
        current_user.full_name = payload["full_name"]
    
    if "email" in payload and payload["email"]:
        current_user.email = payload["email"]

    db.commit()
    db.refresh(current_user)
    return {"message": "Profile updated successfully", "user": {"id": current_user.id, "fullName": current_user.full_name, "email": current_user.email}}


@router.post("/settings/change-password")
def change_admin_password(
    payload: dict,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    from app.core.security import verify_password
    current_pass = payload.get("currentPassword") or payload.get("current_password") or ""
    new_pass = payload.get("newPassword") or payload.get("new_password") or ""

    if not current_pass or not new_pass:
        raise HTTPException(status_code=400, detail="Current password and new password are required")

    if not verify_password(current_pass, current_user.hashed_password):
        raise HTTPException(status_code=400, detail="Current password is incorrect")

    current_user.hashed_password = get_password_hash(new_pass.strip())
    db.commit()
    return {"message": "Password changed successfully"}


# Courses (read-only listing for the admin's own school -- used to let a School
# Admin pick a course when creating a course-scoped announcement; course
# creation/editing itself stays a Teacher responsibility elsewhere in the app).
@router.get("/courses", response_model=List[CourseResponse])
def get_admin_courses(school_id: str = Depends(get_current_school_id), db: Session = Depends(get_db)):
    return db.query(Course).filter(Course.school_id == school_id).all()

# ---------------------------------------------------------------------------
# Formal report card / progress report PDF (Task #49) -- Admin variant:
# any student in the school, no course-ownership restriction (see the
# Teacher-scoped counterpart in teacher.py for teachers).
# ---------------------------------------------------------------------------
@router.get("/students/{student_id}/report-card")
def get_student_report_card_admin(
    student_id: str,
    school_id: str = Depends(get_current_school_id),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    from fastapi.responses import StreamingResponse
    from app.services.report_card import generate_report_card_pdf

    student = db.query(User).filter(
        User.id == student_id, User.school_id == school_id, User.role == UserRole.STUDENT
    ).first()
    if not student:
        raise HTTPException(status_code=404, detail="Student not found")

    buf = generate_report_card_pdf(db, current_user.school, student)
    filename = f"report_card_{student.full_name.replace(' ', '_')}.pdf"
    return StreamingResponse(buf, media_type="application/pdf", headers={"Content-Disposition": f"attachment; filename={filename}"})


# ---------------------------------------------------------------------------
# Search across course content (Task #52) -- Admin variant: every course in
# the school, no ownership restriction.
# ---------------------------------------------------------------------------
@router.get("/search")
def search_course_content_admin(
    q: str = "",
    school_id: str = Depends(get_current_school_id),
    db: Session = Depends(get_db),
):
    from app.services.search_service import search_course_content

    return search_course_content(db, school_id, None, q)


# ---------------------------------------------------------------------------
# Bulk CSV import/export for rostering (Task #53). Students already have a
# bulk .xlsx upload flow (see /students/upload-excel above) -- this adds the
# matching bulk-import path for Teachers (true CSV, since a teacher roster
# is just name+email, no grade/section columns to justify Excel), plus a
# CSV export of the school's current roster for either role.
# ---------------------------------------------------------------------------
@router.get("/teachers/csv-template")
def download_teacher_csv_template():
    import io
    from fastapi.responses import StreamingResponse

    buf = io.StringIO()
    buf.write("Name,Email\n")
    buf.write("Jane Smith,jane.smith@example.com\n")
    buf.write("John Doe,john.doe@example.com\n")
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=teacher_upload_template.csv"},
    )


@router.post("/teachers/upload-csv")
async def upload_teachers_csv(
    file: UploadFile = File(...),
    school_id: str = Depends(get_current_school_id),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """
    Bulk uploads teachers from a .csv file (Name, Email columns), auto-
    generating credentials the same way the student Excel upload does.
    """
    import csv
    import io as _io

    if not file.filename.endswith(".csv"):
        raise HTTPException(status_code=400, detail="Only .csv files are supported")

    contents = (await file.read()).decode("utf-8-sig", errors="replace")
    reader = csv.reader(_io.StringIO(contents))
    rows = list(reader)
    if not rows:
        raise HTTPException(status_code=400, detail="Uploaded CSV file is empty")

    header_row = [c.strip().lower() for c in rows[0]]
    name_idx = next((i for i, c in enumerate(header_row) if "name" in c), -1)
    email_idx = next((i for i, c in enumerate(header_row) if "email" in c), -1)
    if name_idx == -1 or email_idx == -1:
        raise HTTPException(status_code=400, detail="CSV file must contain 'Name' and 'Email' columns")

    existing_emails = {u[0].lower() for u in db.query(User.email).all()}

    created_count = 0
    invalid_count = 0
    errors = []
    credentials_list = []

    for row_num, row in enumerate(rows[1:], start=2):
        if not row or all(not c.strip() for c in row):
            continue
        raw_name = row[name_idx].strip() if name_idx < len(row) else ""
        raw_email = row[email_idx].strip().lower() if email_idx < len(row) else ""

        if not raw_name or not raw_email:
            invalid_count += 1
            errors.append(f"Row {row_num}: Name and Email are both required")
            continue
        if raw_email in existing_emails:
            invalid_count += 1
            errors.append(f"Row {row_num}: Email '{raw_email}' is already in use")
            continue

        raw_password = generate_temp_password(raw_name, fallback="Teacher")
        teacher = User(
            email=raw_email, hashed_password=get_password_hash(raw_password),
            full_name=raw_name, role=UserRole.TEACHER, school_id=school_id,
        )
        db.add(teacher)
        db.commit()
        db.refresh(teacher)
        existing_emails.add(raw_email)

        created_count += 1
        credentials_list.append({
            "id": teacher.id, "name": raw_name, "email": raw_email, "initialPassword": raw_password,
        })

    log = ActivityLog(
        school_id=school_id, user_id=current_user.id, user_name=current_user.full_name,
        action="BULK_TEACHER_UPLOAD", details=f"Bulk imported {created_count} teachers from CSV file '{file.filename}'",
    )
    db.add(log)
    db.commit()

    return {
        "message": f"Successfully onboarded {created_count} teachers.",
        "createdCount": created_count,
        "invalidCount": invalid_count,
        "errors": errors,
        "credentials": credentials_list,
    }


@router.get("/roster/export.csv")
def export_roster_csv(
    role: str = "all",
    school_id: str = Depends(get_current_school_id),
    db: Session = Depends(get_db),
):
    """CSV export of the current roster -- students, teachers, or both."""
    import csv
    import io as _io
    from fastapi.responses import StreamingResponse

    role_norm = role.lower().strip()
    query = db.query(User).filter(User.school_id == school_id)
    if role_norm == "student":
        query = query.filter(User.role == UserRole.STUDENT)
    elif role_norm == "teacher":
        query = query.filter(User.role == UserRole.TEACHER)
    else:
        query = query.filter(User.role.in_([UserRole.STUDENT, UserRole.TEACHER]))
    users = query.order_by(User.role, User.full_name).all()

    buf = _io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["Name", "Email", "Role", "Grade", "Section", "Active"])
    for u in users:
        grade_name = u.grade.name if (u.role == UserRole.STUDENT and u.grade) else ""
        writer.writerow([u.full_name, u.email, u.role.value, grade_name, u.section or "", "Yes" if u.is_active else "No"])

    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename=roster_export_{role_norm}.csv"},
    )


# ---------------------------------------------------------------------------
# SSO scaffolding (Task #62) -- lets a School Admin configure Google/Clever/
# ClassLink SSO for their own school. Rows are created lazily (on first
# save) rather than pre-seeded for every school, but this endpoint always
# reports all three providers so the settings UI has something to render
# even before a school has touched any of them.
# ---------------------------------------------------------------------------
@router.get("/sso-config", response_model=List[SSOConfigResponse])
def get_sso_config(
    school_id: str = Depends(get_current_school_id),
    db: Session = Depends(get_db),
):
    from app.models.sso_config import SSOConfiguration
    from app.services.sso_service import PROVIDERS

    existing = {
        c.provider: c
        for c in db.query(SSOConfiguration).filter(SSOConfiguration.school_id == school_id).all()
    }
    results = []
    for provider, meta in PROVIDERS.items():
        c = existing.get(provider)
        results.append(SSOConfigResponse(
            provider=provider,
            label=meta["label"],
            is_enabled=bool(c and c.is_enabled),
            client_id=c.client_id if c else None,
            has_secret=bool(c and c.client_secret),
            domain_restriction=c.domain_restriction if c else None,
            updated_at=c.updated_at if c else None,
        ))
    return results


@router.put("/sso-config/{provider}", response_model=SSOConfigResponse)
def update_sso_config(
    provider: str,
    payload: SSOConfigUpdate,
    school_id: str = Depends(get_current_school_id),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    from app.models.sso_config import SSOConfiguration
    from app.services.sso_service import PROVIDERS

    if provider not in PROVIDERS:
        raise HTTPException(status_code=404, detail="Unknown SSO provider")

    config = db.query(SSOConfiguration).filter(
        SSOConfiguration.school_id == school_id, SSOConfiguration.provider == provider
    ).first()
    if not config:
        config = SSOConfiguration(school_id=school_id, provider=provider)
        db.add(config)

    if payload.is_enabled is not None:
        config.is_enabled = payload.is_enabled
    if payload.client_id is not None:
        config.client_id = payload.client_id or None
    if payload.client_secret is not None:
        # Empty string clears a stored secret; omitted (None) leaves it untouched.
        config.client_secret = payload.client_secret or None
    if payload.domain_restriction is not None:
        config.domain_restriction = (payload.domain_restriction or "").strip().lower() or None

    config.updated_by_id = current_user.id
    db.add(ActivityLog(
        school_id=school_id, user_id=current_user.id, user_name=current_user.full_name,
        action="SSO_CONFIG_UPDATED", details=f"{PROVIDERS[provider]['label']} SSO settings updated",
    ))
    db.commit()
    db.refresh(config)

    return SSOConfigResponse(
        provider=provider,
        label=PROVIDERS[provider]["label"],
        is_enabled=config.is_enabled,
        client_id=config.client_id,
        has_secret=bool(config.client_secret),
        domain_restriction=config.domain_restriction,
        updated_at=config.updated_at,
    )
@router.post("/complaints", response_model=ComplaintResponse)
def create_complaint(
    payload: ComplaintCreate,
    db: Session = Depends(get_db),
    school_id: str = Depends(get_current_school_id),
    current_user: User = Depends(get_current_user),
):
    """Task #66: lets a School Admin raise a complaint/support ticket that
    the Super Admin sees in their own queue (see super_admin.py)."""
    subject = (payload.subject or "").strip()
    description = (payload.description or "").strip()
    if not subject or not description:
        raise HTTPException(status_code=400, detail="Subject and description are required")

    complaint = Complaint(
        school_id=school_id,
        submitted_by_id=current_user.id,
        submitted_by_name=current_user.full_name,
        subject=subject,
        description=description,
    )
    db.add(complaint)
    db.add(ActivityLog(
        school_id=school_id, user_id=current_user.id, user_name=current_user.full_name,
        action="COMPLAINT_SUBMITTED", details=subject,
    ))
    db.commit()
    db.refresh(complaint)
    complaint.school_name = complaint.school.name if complaint.school else None
    return complaint


@router.get("/complaints", response_model=List[ComplaintResponse])
def list_my_complaints(
    db: Session = Depends(get_db),
    school_id: str = Depends(get_current_school_id),
):
    complaints = (
        db.query(Complaint)
        .filter(Complaint.school_id == school_id)
        .order_by(Complaint.created_at.desc())
        .all()
    )
    for c in complaints:
        c.school_name = c.school.name if c.school else None
    return complaints
